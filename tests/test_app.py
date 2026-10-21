import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate
from werkzeug.security import generate_password_hash

from app import PageNumberCanvas, create_app, get_db


class RepairReportAppTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp_dir.name) / "test.sqlite3"
        self.app = create_app(
            {
                "TESTING": True,
                "DATABASE": str(self.database_path),
                "UPLOAD_DIR": str(Path(self.temp_dir.name) / "uploads"),
                "SECRET_KEY": "test",
            }
        )
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pdf_page_numbers_include_total_pages(self):
        page_numbers = []

        class RecordingPageNumberCanvas(PageNumberCanvas):
            def drawCentredString(self, x, y, text):
                page_numbers.append(text)
                super().drawCentredString(x, y, text)

        output = io.BytesIO()
        paragraph_style = getSampleStyleSheet()["Normal"]
        SimpleDocTemplate(output, pagesize=A4).build(
            [
                Paragraph("Pierwsza strona", paragraph_style),
                PageBreak(),
                Paragraph("Druga strona", paragraph_style),
            ],
            canvasmaker=RecordingPageNumberCanvas,
        )

        self.assertEqual(page_numbers, ["1/2", "2/2"])

    def test_create_record_calculates_downtime_and_cost(self):
        response = self.client.post(
            "/",
            data={
                "occurrence_at": "2026-08-30T08:00",
                "repaired_at": "2026-08-30T10:30",
                "machine_name": "Prasa 1",
                "machine_type": "Hydrauliczna",
                "serial_number": "H-001",
                "event_type": "Awaria",
                "description": "Wymiana czujnika",
                "cause": "Zużycie",
                "repair_method": "Wymiana",
                "replaced_parts": "Czujnik",
                "labor_cost": "1",
                "material_cost": "20",
                "spare_parts_cost": "100",
                "service_cost": "0",
            },
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Rekord został zapisany.", page)
        self.assertIn("2.5", page)
        self.assertIn("200.0", page)

    def test_guest_can_view_and_export_but_cannot_create_records(self):
        app = create_app(
            {
                "TESTING": False,
                "REQUIRE_LOGIN": False,
                "DATABASE": str(self.database_path),
                "UPLOAD_DIR": str(Path(self.temp_dir.name) / "uploads"),
                "SECRET_KEY": "test",
            }
        )
        client = app.test_client()

        self.assertEqual(client.get("/").status_code, 200)
        self.assertEqual(client.get("/export").status_code, 200)

        response = client.post(
            "/",
            data={
                "occurrence_at": "2026-08-30T08:00",
                "machine_name": "Maszyna gościa",
                "event_type": "Awaria",
                "description": "Nieautoryzowany wpis",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Brak uprawnień do dodawania rekordów.", response.get_data(as_text=True))

        with app.app_context():
            record_count = get_db().execute("SELECT COUNT(*) FROM repair_records").fetchone()[0]
        self.assertEqual(record_count, 0)

    def test_polish_date_format_is_displayed_and_accepted(self):
        response = self.client.post(
            "/",
            data={
                "occurrence_at": "08-10-2026 14:35",
                "repaired_at": "08-10-2026 15:05",
                "machine_name": "Maszyna dat",
                "machine_type": "Testowa",
                "serial_number": "DATE-001",
                "event_type": "Awaria",
                "description": "Test formatu daty",
            },
        )
        self.assertEqual(response.status_code, 302)

        page = self.client.get("/").get_data(as_text=True)
        self.assertIn("08-10-2026\n14:35", page)
        for path in ("/new", "/", "/filter"):
            with self.subTest(path=path):
                page = self.client.get(path).get_data(as_text=True)
                self.assertIn('type="datetime-local"', page)
                self.assertIn("syncDisplayDateTime(this)", page)
                self.assertIn("syncNativeDateTime(this)", page)
                self.assertIn('onclick="openDateTimePicker(this)"', page)
                self.assertIn("picker.showPicker()", page)

        with self.app.app_context():
            record_id = get_db().execute(
                """
                SELECT r.id
                FROM repair_records r
                JOIN machines m ON m.id = r.machine_id
                WHERE m.serial_number = ?
                """,
                ("DATE-001",),
            ).fetchone()
            self.assertIsNotNone(record_id)
            record_id = int(record_id["id"])

        record = self.client.get(f"/api/record/{record_id}").get_json()
        self.assertEqual(record["occurrence_at"], "2026-10-08 14:35")
        history = self.client.get(
            "/api/machines?machine=Maszyna%20dat&machine_type=Testowa&serial_number=DATE-001"
        ).get_json()
        self.assertEqual(history["previous_entries"][0]["occurrence_at"], "08-10-2026 14:35")

        filtered_page = self.client.get("/filter?occurrence_from=08-10-2026").get_data(as_text=True)
        self.assertIn("Test formatu daty", filtered_page)
        self.assertIn('value="08-10-2026"', filtered_page)

    def test_machine_api_and_filtered_export(self):
        self.client.post(
            "/",
            data={
                "occurrence_at": "2026-08-30T08:00",
                "repaired_at": "2026-08-30T09:00",
                "machine_name": "Linia A",
                "machine_type": "Pakowarka",
                "serial_number": "P-100",
                "event_type": "Przegląd",
                "description": "Kontrola napędu",
            },
        )
        self.client.post(
            "/",
            data={
                "occurrence_at": "2026-08-31T08:00",
                "repaired_at": "2026-08-31T09:00",
                "machine_name": "Linia B",
                "machine_type": "Pakowarka",
                "serial_number": "P-200",
                "event_type": "Awaria",
                "description": "Naprawa silnika",
                "notes": "Kontrola łożyska",
                "participants": "Jan Kowalski, Anna Nowak",
            },
        )

        api_response = self.client.get("/api/machines?machine=Linia%20A&machine_type=Pakowarka")
        payload = api_response.get_json()
        self.assertEqual(payload["types"], ["Pakowarka"])
        self.assertEqual(payload["serial_numbers"], ["P-100"])
        self.assertEqual(len(payload["previous_entries"]), 1)

        export_response = self.client.get("/export?machine=Linia%20B&description_terms=silnika")
        self.assertEqual(export_response.status_code, 200)
        workbook = load_workbook(io.BytesIO(export_response.data))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][4], "Linia B")
        self.assertEqual(rows[1][13], "31-08-2026 08:00")
        self.assertEqual(rows[1][14], "31-08-2026 09:00")
        self.assertEqual(rows[0][12], "Osoby biorące udział w interwencji")
        self.assertEqual(rows[1][12], "Jan Kowalski, Anna Nowak")

        filter_response = self.client.get("/filter?description_terms=łożyska")
        filter_page = filter_response.get_data(as_text=True)
        self.assertIn("Linia B", filter_page)
        self.assertIn("31-08-2026\n08:00", filter_page)

        notes_export_response = self.client.get("/export?description_terms=łożyska")
        self.assertEqual(notes_export_response.status_code, 200)
        notes_workbook = load_workbook(io.BytesIO(notes_export_response.data))
        notes_rows = list(notes_workbook.active.iter_rows(values_only=True))
        self.assertEqual(len(notes_rows), 2)
        self.assertEqual(notes_rows[1][4], "Linia B")

    def test_admin_can_edit_machine_details_for_all_linked_records(self):
        for day in ("30", "31"):
            self.client.post(
                "/",
                data={
                    "occurrence_at": f"2026-08-{day}T08:00",
                    "machine_name": "Linia A",
                    "machine_type": "Pakowarka",
                    "serial_number": "P-100",
                    "event_type": "Awaria",
                    "description": f"Wpis {day}",
                },
            )
        self.client.post(
            "/",
            data={
                "occurrence_at": "2026-09-01T08:00",
                "machine_name": "Linia B",
                "machine_type": "Pakowarka",
                "serial_number": "P-200",
                "event_type": "Awaria",
                "description": "Druga maszyna",
            },
        )
        with self.client.session_transaction() as session:
            session["logged_in"] = True
            session["role"] = "admin"

        with self.app.app_context():
            machine_id = get_db().execute(
                "SELECT id FROM machines WHERE serial_number = ?", ("P-100",)
            ).fetchone()["id"]

        response = self.client.post(
            f"/admin/machines/{machine_id}/edit",
            data={
                "name": "Linia A po zmianie",
                "machine_type": "Nowa pakowarka",
                "serial_number": "P-101",
            },
        )
        self.assertEqual(response.status_code, 302)
        page = self.client.get("/admin/machines").get_data(as_text=True)
        self.assertIn("Linia A po zmianie", page)
        self.assertIn("P-101", page)
        with self.app.app_context():
            names = get_db().execute(
                """
                SELECT DISTINCT m.name, m.machine_type, m.serial_number
                FROM repair_records r JOIN machines m ON m.id = r.machine_id
                WHERE r.description IN (?, ?)
                """,
                ("Wpis 30", "Wpis 31"),
            ).fetchall()
        self.assertEqual(len(names), 1)
        self.assertEqual(tuple(names[0]), ("Linia A po zmianie", "Nowa pakowarka", "P-101"))

        response = self.client.post(
            f"/admin/machines/{machine_id}/edit",
            data={
                "name": "Konflikt",
                "machine_type": "Pakowarka",
                "serial_number": "P-200",
            },
            follow_redirects=True,
        )
        self.assertIn("Maszyna o takich danych lub numerze seryjnym już istnieje.", response.get_data(as_text=True))
        with self.app.app_context():
            serial_number = get_db().execute(
                "SELECT serial_number FROM machines WHERE id = ?", (machine_id,)
            ).fetchone()["serial_number"]
        self.assertEqual(serial_number, "P-101")

    def test_admin_can_delete_machine_with_records_and_attachments(self):
        self.client.post(
            "/",
            data={
                "occurrence_at": "2026-08-30T08:00",
                "machine_name": "Linia do usunięcia",
                "machine_type": "Pakowarka",
                "serial_number": "P-DEL",
                "event_type": "Awaria",
                "description": "Wpis do usunięcia",
            },
        )
        with self.app.app_context():
            record_id = int(
                get_db().execute(
                    "SELECT id FROM repair_records WHERE description = ?",
                    ("Wpis do usunięcia",),
                ).fetchone()["id"]
            )
        attachment_response = self.client.post(
            f"/api/record/{record_id}/attachments",
            data={"attachments": (io.BytesIO(b"test file"), "test.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(attachment_response.status_code, 200)
        attachment_dir = Path(self.app.config["UPLOAD_DIR"]) / str(record_id)
        self.assertTrue(attachment_dir.exists())

        with self.client.session_transaction() as session:
            session["logged_in"] = True
            session["role"] = "admin"
        with self.app.app_context():
            machine_id = get_db().execute(
                "SELECT id FROM machines WHERE serial_number = ?", ("P-DEL",)
            ).fetchone()["id"]

        response = self.client.post(f"/admin/machines/{machine_id}/delete")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(attachment_dir.exists())
        with self.app.app_context():
            self.assertEqual(
                get_db().execute(
                    "SELECT COUNT(*) FROM repair_records WHERE machine_id = ?", (machine_id,)
                ).fetchone()[0],
                0,
            )
            self.assertEqual(
                get_db().execute(
                    "SELECT COUNT(*) FROM attachments WHERE record_id = ?", (record_id,)
                ).fetchone()[0],
                0,
            )
            self.assertIsNone(
                get_db().execute("SELECT id FROM machines WHERE id = ?", (machine_id,)).fetchone()
            )

    def test_machine_management_requires_admin_role(self):
        response = self.client.get("/admin/machines")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/")
        edit_response = self.client.post(
            "/admin/machines/1/edit",
            data={"name": "Nieuprawniona zmiana", "machine_type": "Typ", "serial_number": "123"},
        )
        delete_response = self.client.post("/admin/machines/1/delete")
        self.assertEqual(edit_response.status_code, 302)
        self.assertEqual(delete_response.status_code, 302)
        self.assertEqual(edit_response.location, "/")
        self.assertEqual(delete_response.location, "/")

    def test_serial_number_reuses_existing_machine_even_with_different_label(self):
        first = {
            "occurrence_at": "2026-08-30T08:00",
            "machine_name": "BOSCHERT",
            "machine_type": "Twin 1000",
            "serial_number": "282",
            "event_type": "Awaria",
            "description": "Pierwszy wpis",
        }
        second = {
            **first,
            "machine_name": "Inna nazwa",
            "machine_type": "Inny typ",
            "description": "Drugi wpis",
        }

        self.client.post("/", data=first)
        self.client.post("/", data=second)

        with self.app.app_context():
            machines = get_db().execute(
                "SELECT id, name, machine_type FROM machines WHERE serial_number = ?",
                ("282",),
            ).fetchall()
            records = get_db().execute(
                "SELECT DISTINCT machine_id FROM repair_records WHERE description IN (?, ?)",
                ("Pierwszy wpis", "Drugi wpis"),
            ).fetchall()

        self.assertEqual(len(machines), 1)
        self.assertEqual(machines[0]["name"], "BOSCHERT")
        self.assertEqual(machines[0]["machine_type"], "Twin 1000")
        self.assertEqual(len(records), 1)

    def test_recent_records_limit_all_displays_all_rows(self):
        valid_dates = [
            ("2026-08-30T08:00", "2026-08-30T09:00"),
            ("2026-08-31T08:00", "2026-08-31T09:00"),
            ("2026-09-01T08:00", "2026-09-01T09:00"),
        ]
        for idx, (occurrence_at, repaired_at) in enumerate(valid_dates):
            self.client.post(
                "/",
                data={
                    "occurrence_at": occurrence_at,
                    "repaired_at": repaired_at,
                    "machine_name": f"Maszyna {idx}",
                    "machine_type": "Typ",
                    "serial_number": f"SER-{idx}",
                    "event_type": "Awaria",
                    "description": f"Opis {idx}",
                },
            )

        response = self.client.get("/?limit=all")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Ostatnie wszystkie rekordów", page)
        self.assertIn("Opis 0", page)
        self.assertIn("Opis 1", page)
        self.assertIn("Opis 2", page)

    def test_record_summaries_are_limited_in_recent_and_filtered_tables(self):
        description = "Opis z limitem\n" + ("szczegółowy tekst " * 60)
        self.client.post(
            "/",
            data={
                "occurrence_at": "2026-08-30T08:00",
                "repaired_at": "2026-08-30T09:00",
                "machine_name": "Maszyna testowa",
                "machine_type": "Typ",
                "serial_number": "SER-400",
                "event_type": "Awaria",
                "description": description,
            },
        )
        expected_summary = description[:397] + "..."

        for path in ("/?limit=all", "/filter?description_terms=Opis"):
            with self.subTest(path=path):
                response = self.client.get(path)
                page = response.get_data(as_text=True)
                self.assertEqual(response.status_code, 200)
                self.assertIn(expected_summary, page)
                self.assertNotIn(description, page)
                self.assertIn("class=\"record-summary\"", page)

    def test_records_without_repair_date_are_highlighted_in_recent_and_filtered_tables(self):
        self.client.post(
            "/",
            data={
                "occurrence_at": "2026-08-30T08:00",
                "machine_name": "Maszyna bez daty naprawienia",
                "machine_type": "Typ",
                "serial_number": "SER-401",
                "event_type": "Awaria",
                "description": "Oczekuje na naprawę",
            },
        )

        for path in ("/?limit=all", "/filter?description_terms=Oczekuje"):
            with self.subTest(path=path):
                response = self.client.get(path)
                page = response.get_data(as_text=True)
                self.assertEqual(response.status_code, 200)
                self.assertIn('class="date-no-repair"', page)
                self.assertIn('title="Brak daty naprawienia"', page)

    def test_import_excel_file(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Maszyna",
                "Typ",
                "Numer",
                "Rodzaj zdarzenia",
                "Opis",
                "Osoby biorące udział w interwencji",
                "Data wystąpienia",
                "Data naprawienia/ interwencji",
            ]
        )
        sheet.append(
            [
                "Gilotyna",
                "Tnąca",
                "G-12",
                "Awaria",
                "Wymiana noża",
                "Marek Nowak, Piotr Kowalski",
                "2026-08-20 07:00",
                "2026-08-20 08:30",
            ]
        )
        sheet.append(
            [
                "",
                "",
                "",
                "Awaria",
                "Niekompletny wiersz",
                "",
                "2026-08-21 07:00",
                "2026-08-21 08:00",
            ]
        )
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = self.client.post(
            "/import",
            data={"file": (buffer, "raport.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Zaimportowano 1 rekordów.", page)
        self.assertIn("Gilotyna", page)
        with self.app.app_context():
            participants = get_db().execute(
                "SELECT participants FROM repair_records WHERE description = ?",
                ("Wymiana noża",),
            ).fetchone()["participants"]
        self.assertEqual(participants, "Marek Nowak, Piotr Kowalski")

    def test_user_login_and_create_permission(self):
        app = create_app(
            {
                "TESTING": False,
                "REQUIRE_LOGIN": True,
                "DATABASE": str(self.database_path),
                "SECRET_KEY": "test",
            }
        )
        client = app.test_client()

        with app.app_context():
            get_db().execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT NOT NULL UNIQUE,
                    password_hash TEXT NOT NULL,
                    is_admin INTEGER NOT NULL DEFAULT 0,
                    can_view_records INTEGER NOT NULL DEFAULT 1,
                    can_create_records INTEGER NOT NULL DEFAULT 0,
                    can_edit_records INTEGER NOT NULL DEFAULT 0,
                    can_delete_records INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )
            get_db().execute(
                "INSERT INTO users(username, password_hash, is_admin, can_view_records, can_create_records, can_edit_records, can_delete_records, created_at, updated_at) VALUES (?, ?, 0, 1, 1, 1, 0, ?, ?)",
                ("operator", generate_password_hash("test"), "2026-01-01T00:00:00", "2026-01-01T00:00:00"),
            )
            get_db().commit()

        response = client.post(
            "/login",
            data={"login": "operator", "password": "test"},
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Strona główna", response.get_data(as_text=True))
        self.assertIn("Zalogowany: <strong>operator</strong>", response.get_data(as_text=True))

        response = client.get("/new")
        self.assertEqual(response.status_code, 200)

        response = client.post(
            "/new",
            data={
                "occurrence_at": "2026-09-21T08:00",
                "machine_name": "Maszyna operatora",
                "machine_type": "Testowa",
                "serial_number": "OP-001",
                "event_type": "Awaria",
                "description": "Rekord utworzony przez operatora",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)

        with app.app_context():
            activity = get_db().execute(
                """
                SELECT action, username, target_type
                FROM user_activity_log
                WHERE action = 'create_record'
                ORDER BY id DESC
                LIMIT 1
                """
            ).fetchone()
        self.assertIsNotNone(activity)
        self.assertEqual(activity["action"], "create_record")
        self.assertEqual(activity["username"], "operator")
        self.assertEqual(activity["target_type"], "record")

        with app.app_context():
            get_db().execute("DELETE FROM users WHERE username = ?", ("operator",))
            get_db().commit()

    def test_admin_can_edit_existing_user_permissions(self):
        with self.app.app_context():
            cursor = get_db().execute(
                """
                INSERT INTO users(
                    username, password_hash, is_admin, can_view_records,
                    can_create_records, can_edit_records, can_delete_records,
                    created_at, updated_at
                ) VALUES (?, ?, 0, 1, 1, 0, 0, ?, ?)
                """,
                (
                    "permissions_user",
                    generate_password_hash("unchanged"),
                    "2026-01-01T00:00:00",
                    "2026-01-01T00:00:00",
                ),
            )
            user_id = int(cursor.lastrowid)
            get_db().commit()

        with self.client.session_transaction() as session:
            session["logged_in"] = True
            session["role"] = "admin"
            session["user_id"] = 1
            session["username"] = "admin"

        response = self.client.get(f"/admin/users?edit={user_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn('name="user_id"', response.get_data(as_text=True))
        self.assertIn("permissions_user", response.get_data(as_text=True))

        response = self.client.post(
            "/admin/users",
            data={
                "user_id": str(user_id),
                "username": "permissions_user",
                "is_admin": "0",
                "can_view_records": "on",
                "can_create_records": "on",
                "can_edit_records": "on",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)

        with self.app.app_context():
            user = get_db().execute(
                "SELECT can_view_records, can_create_records, can_edit_records, can_delete_records FROM users WHERE id = ?",
                (user_id,),
            ).fetchone()
        self.assertEqual(tuple(user), (1, 1, 1, 0))

    def test_admin_can_export_full_activity_log(self):
        with self.app.app_context():
            get_db().execute(
                """
                INSERT INTO user_activity_log(
                    user_id, username, action, target_type, target_id, details, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (1, "admin", "create_record", "record", 123, "Testowe zdarzenie", "2026-09-21T11:50:00+00:00"),
            )
            get_db().commit()

        with self.client.session_transaction() as session:
            session["logged_in"] = True
            session["role"] = "admin"
            session["user_id"] = 1
            session["username"] = "admin"

        response = self.client.get("/admin/users/activity/export")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(io.BytesIO(response.data))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Data", "Użytkownik", "Akcja", "Typ obiektu", "ID obiektu", "Szczegóły"))
        self.assertIn(("21-09-2026 13:50", "admin", "create_record", "record", 123, "Testowe zdarzenie"), rows)

    def test_editor_can_open_edit_controls_in_filter_and_machine_history(self):
        app = create_app(
            {
                "TESTING": False,
                "REQUIRE_LOGIN": True,
                "DATABASE": str(self.database_path),
                "SECRET_KEY": "test",
            }
        )
        client = app.test_client()
        with app.app_context():
            get_db().execute(
                """
                INSERT INTO users(
                    username, password_hash, is_admin, can_view_records,
                    can_create_records, can_edit_records, can_delete_records,
                    created_at, updated_at
                ) VALUES (?, ?, 0, 1, 1, 1, 1, ?, ?)
                """,
                (
                    "editor",
                    generate_password_hash("editor-pass"),
                    "2026-01-01T00:00:00",
                    "2026-01-01T00:00:00",
                ),
            )
            get_db().commit()

        response = client.post(
            "/login",
            data={"login": "editor", "password": "editor-pass"},
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)

        filter_page = client.get("/filter")
        new_record_page = client.get("/new")
        for page in (filter_page, new_record_page):
            html = page.get_data(as_text=True)
            self.assertEqual(page.status_code, 200)
            self.assertIn("const CAN_EDIT = 1;", html)
            self.assertIn("(USER_ROLE === 'admin' || CAN_EDIT)", html)


if __name__ == "__main__":
    unittest.main()
